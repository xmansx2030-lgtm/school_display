"""School-manager schedule mutations must be scoped, validated and atomic."""

from datetime import time
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import School, SubscriptionPlan, UserProfile
from dashboard.excel_import import apply_import, build_template_bytes, parse_workbook, undo_import
from schedule.models import Break, ClassLesson, DaySchedule, Period, SchoolClass, SchoolSettings, Subject, Teacher
from subscriptions.models import SchoolSubscription


class ScheduleMutationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.school = School.objects.create(name="مدرسة سلامة الجدول", slug="schedule-safety")
        cls.other_school = School.objects.create(name="مدرسة أخرى", slug="schedule-other")
        cls.settings = SchoolSettings.objects.create(school=cls.school, name=cls.school.name)
        cls.other_settings = SchoolSettings.objects.create(school=cls.other_school, name=cls.other_school.name)
        cls.user = get_user_model().objects.create_user(username="schedule-safety-manager")
        profile = UserProfile.objects.create(user=cls.user, active_school=cls.school)
        profile.schools.add(cls.school)
        plan = SubscriptionPlan.objects.create(code="schedule-safety", name="اختبار", price=10, duration_days=365)
        SchoolSubscription.objects.create(school=cls.school, plan=plan, starts_at=timezone.localdate(), status="active")
        cls.day = DaySchedule.objects.create(settings=cls.settings, weekday=7, periods_count=2)
        cls.period1 = Period.objects.create(day=cls.day, index=1, starts_at=time(7), ends_at=time(7, 45))
        cls.period2 = Period.objects.create(day=cls.day, index=2, starts_at=time(8), ends_at=time(8, 45))
        cls.school_class = SchoolClass.objects.create(settings=cls.settings, name="الأول أ")
        cls.other_class = SchoolClass.objects.create(settings=cls.other_settings, name="فصل آخر")
        cls.teacher = Teacher.objects.create(school=cls.school, name="معلم المدرسة")
        cls.subject = Subject.objects.create(school=cls.school, name="رياضيات")
        cls.other_teacher = Teacher.objects.create(school=cls.other_school, name="معلم مدرسة أخرى")
        cls.other_subject = Subject.objects.create(school=cls.other_school, name="مادة مدرسة أخرى")
        cls.lesson = ClassLesson.objects.create(settings=cls.settings, school_class=cls.school_class,
            weekday=7, period_index=1, teacher=cls.teacher, subject=cls.subject)

    def setUp(self):
        cache.clear()
        self.client.force_login(self.user)

    def autofill(self, **overrides):
        values = {"start_time": "07:00", "target_periods_count": "3", "period_minutes": "45",
                  "gap_minutes": "5", "break_after": "1", "break_minutes": "20"}
        values.update(overrides)
        return self.client.post(reverse("dashboard:day_autofill", args=[7]), values)

    def assert_original_day(self):
        self.day.refresh_from_db()
        self.assertEqual(self.day.periods_count, 2)
        self.assertEqual(list(self.day.periods.order_by("index").values_list("id", "starts_at", "ends_at")),
                         [(self.period1.pk, time(7), time(7, 45)), (self.period2.pk, time(8), time(8, 45))])

    def test_invalid_autofill_input_does_not_change_declared_count(self):
        self.assertEqual(self.autofill(start_time="bad").status_code, 302)
        self.assert_original_day()

    def test_autofill_rejects_schedule_crossing_midnight(self):
        self.assertEqual(self.autofill(start_time="23:00", period_minutes="90").status_code, 302)
        self.assert_original_day()

    def test_autofill_failure_rolls_back_deleted_periods_and_partial_replacement(self):
        save = Period.save
        created = 0

        def fail_second(instance, *args, **kwargs):
            nonlocal created
            created += 1
            if created == 2:
                raise RuntimeError("simulated interruption while saving")
            return save(instance, *args, **kwargs)

        with patch.object(Period, "save", fail_second):
            self.assertEqual(self.autofill().status_code, 302)
        self.assert_original_day()

    def test_reindex_can_swap_existing_unique_indexes(self):
        Period.objects.filter(pk=self.period1.pk).update(starts_at=time(9), ends_at=time(9, 45))
        response = self.client.post(reverse("dashboard:day_reindex", args=[7]))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(list(self.day.periods.order_by("starts_at").values_list("index", flat=True)), [1, 2])

    def test_timetable_post_cannot_fall_back_from_foreign_class_to_first_local_class(self):
        response = self.client.post(reverse("dashboard:timetable_day"), {"weekday": 7, "class_id": self.other_class.pk})
        self.assertIn(response.status_code, (400, 404))
        self.assertTrue(ClassLesson.objects.filter(pk=self.lesson.pk).exists())

    def test_missing_or_invalid_weekday_cannot_write_to_a_default_day(self):
        for weekday in (None, "", "0", "not-a-day"):
            with self.subTest(weekday=weekday):
                data = {"class_id": self.school_class.pk,
                    f"subject-{self.school_class.pk}-1": "", f"teacher-{self.school_class.pk}-1": ""}
                if weekday is not None:
                    data["weekday"] = weekday
                response = self.client.post(reverse("dashboard:timetable_day"), data)
                self.assertEqual(response.status_code, 400)
                self.assertTrue(ClassLesson.objects.filter(pk=self.lesson.pk).exists())

    def test_timetable_rejects_foreign_choices_instead_of_treating_them_as_deletion(self):
        response = self.client.post(reverse("dashboard:timetable_day"), {
            "weekday": 7, "class_id": self.school_class.pk,
            f"subject-{self.school_class.pk}-1": self.other_subject.pk,
            f"teacher-{self.school_class.pk}-1": self.other_teacher.pk,
        }, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.status_code, 400)
        self.lesson.refresh_from_db()
        self.assertEqual(self.lesson.teacher_id, self.teacher.pk)

    def test_partial_timetable_submission_does_not_delete_unsubmitted_rows(self):
        response = self.client.post(reverse("dashboard:timetable_day"), {
            "weekday": 7, "class_id": self.school_class.pk,
            f"subject-{self.school_class.pk}-2": self.subject.pk,
            f"teacher-{self.school_class.pk}-2": self.teacher.pk,
        }, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(ClassLesson.objects.filter(pk=self.lesson.pk).exists())
        self.assertTrue(ClassLesson.objects.filter(settings=self.settings, period_index=2).exists())

    def test_invalid_later_timetable_row_prevents_all_earlier_changes(self):
        response = self.client.post(reverse("dashboard:timetable_day"), {
            "weekday": 7, "class_id": self.school_class.pk,
            f"subject-{self.school_class.pk}-1": "", f"teacher-{self.school_class.pk}-1": "",
            f"subject-{self.school_class.pk}-2": self.subject.pk, f"teacher-{self.school_class.pk}-2": "",
        }, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.status_code, 400)
        self.assertTrue(ClassLesson.objects.filter(pk=self.lesson.pk).exists())
        self.assertFalse(ClassLesson.objects.filter(settings=self.settings, period_index=2).exists())

    def test_explicit_empty_timetable_row_deletes_only_that_row(self):
        second = ClassLesson.objects.create(settings=self.settings, school_class=self.school_class,
            weekday=7, period_index=2, teacher=self.teacher, subject=self.subject)
        response = self.client.post(reverse("dashboard:timetable_day"), {
            "weekday": 7, "class_id": self.school_class.pk,
            f"subject-{self.school_class.pk}-1": "", f"teacher-{self.school_class.pk}-1": "",
        }, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.status_code, 200)
        self.assertFalse(ClassLesson.objects.filter(pk=self.lesson.pk).exists())
        self.assertTrue(ClassLesson.objects.filter(pk=second.pk).exists())

    def test_duplicate_lesson_returns_form_error_instead_of_integrity_error(self):
        response = self.client.post(reverse("dashboard:add_lesson"), {
            "school_class": self.school_class.pk, "weekday": 7, "period_index": 1,
            "subject": self.subject.pk, "teacher": self.teacher.pk, "is_active": "on",
        })
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["form"].errors)
        self.assertEqual(ClassLesson.objects.filter(settings=self.settings).count(), 1)

    def day_form(self):
        return {"periods_count": "2", "is_active": "on", "p-TOTAL_FORMS": "2", "p-INITIAL_FORMS": "2",
                "p-0-id": self.period1.pk, "p-0-index": 1, "p-0-starts_at": "07:00", "p-0-ends_at": "07:45",
                "p-1-id": self.period2.pk, "p-1-index": 2, "p-1-starts_at": "08:00", "p-1-ends_at": "08:45",
                "b-TOTAL_FORMS": "2", "b-INITIAL_FORMS": "0",
                "b-0-label": "فسحة", "b-0-starts_at": "09:00", "b-0-duration_min": "20",
                "b-1-label": "فسحة ثانية", "b-1-starts_at": "09:10", "b-1-duration_min": "20"}

    def test_day_editor_rejects_two_overlapping_breaks(self):
        response = self.client.post(reverse("dashboard:day_edit", args=[7]), self.day_form())
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Break.objects.filter(day=self.day).exists())

    def test_malformed_break_management_data_returns_validation_error(self):
        data = self.day_form()
        data["b-TOTAL_FORMS"] = "invalid"
        response = self.client.post(reverse("dashboard:day_edit", args=[7]), data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Break.objects.filter(day=self.day).exists())

    def test_foreign_lesson_is_not_readable_editable_or_deletable(self):
        other = ClassLesson.objects.create(settings=self.other_settings, school_class=self.other_class,
            weekday=7, period_index=1, teacher=self.other_teacher, subject=self.other_subject)
        self.assertEqual(self.client.get(reverse("dashboard:edit_lesson", args=[other.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse("dashboard:edit_lesson", args=[other.pk]), {}).status_code, 404)
        self.assertEqual(self.client.post(reverse("dashboard:delete_lesson", args=[other.pk])).status_code, 404)
        self.assertTrue(ClassLesson.objects.filter(pk=other.pk).exists())

    def test_valid_autofill_replaces_only_the_active_schools_day(self):
        other = DaySchedule.objects.create(settings=self.other_settings, weekday=7)
        untouched = Period.objects.create(day=other, index=1, starts_at=time(7), ends_at=time(7, 30))
        self.autofill()
        self.day.refresh_from_db()
        self.assertEqual(self.day.periods_count, 3)
        self.assertEqual(self.day.periods.count(), 3)
        self.assertEqual(self.day.breaks.count(), 1)
        self.assertTrue(Period.objects.filter(pk=untouched.pk).exists())

    def test_school_data_batch_rejects_overlong_names_before_creating_anything(self):
        for route, model in (("add_class", SchoolClass), ("add_subject", Subject), ("add_teacher", Teacher)):
            with self.subTest(route=route):
                response = self.client.post(reverse("dashboard:" + route), {"name": "اسم جديد صحيح\n" + "س" * (model._meta.get_field("name").max_length + 1)})
                self.assertEqual(response.status_code, 302)
                self.assertFalse(model.objects.filter(name="اسم جديد صحيح").exists())


class ExcelScheduleSafetyTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(name="مدرسة الاستيراد", slug="import-safety")

    def workbook(self, rows):
        workbook = load_workbook(BytesIO(build_template_bytes()))
        worksheet = workbook["الجدول"]
        worksheet.delete_rows(2, worksheet.max_row)
        for row in rows:
            worksheet.append(row)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    def test_preview_rejects_overlapping_period_times(self):
        parsed = parse_workbook(self.workbook([
            ["الأحد", 1, "07:00", "07:45", "أول أ", "رياضيات", "معلم أ"],
            ["الأحد", 2, "07:30", "08:15", "أول أ", "رياضيات", "معلم أ"],
        ]))
        self.assertTrue(parsed["errors"])

    def test_undo_refuses_to_overwrite_manual_changes_after_import(self):
        parsed = parse_workbook(self.workbook([["الأحد", 1, "07:00", "07:45", "أول أ", "رياضيات", "معلم أ"]]))
        result = apply_import(school=self.school, parsed=parsed)
        lesson = ClassLesson.objects.get(settings__school=self.school)
        replacement = Teacher.objects.create(school=self.school, name="تعديل يدوي لاحق")
        lesson.teacher = replacement
        lesson.save()
        with self.assertRaises(ValueError):
            undo_import(school=self.school, undo=result["undo"])
        lesson.refresh_from_db()
        self.assertEqual(lesson.teacher_id, replacement.pk)

    def test_import_rolls_back_when_it_conflicts_with_existing_break(self):
        settings = SchoolSettings.objects.create(school=self.school, name=self.school.name)
        day = DaySchedule.objects.create(settings=settings, weekday=7)
        Break.objects.create(day=day, starts_at=time(7, 30), duration_min=15)
        parsed = parse_workbook(self.workbook([["الأحد", 1, "07:00", "07:45", "أول أ", "رياضيات", "معلم أ"]]))
        self.assertFalse(parsed["errors"])
        with self.assertRaises(ValueError):
            apply_import(school=self.school, parsed=parsed)
        self.assertFalse(ClassLesson.objects.filter(settings=settings).exists())
        self.assertFalse(Period.objects.filter(day=day).exists())
        self.assertFalse(Teacher.objects.filter(school=self.school).exists())
        self.assertEqual(Break.objects.filter(day=day).count(), 1)

    def test_preview_rejects_fractional_period_numbers_and_overlong_names(self):
        parsed = parse_workbook(self.workbook([["الأحد", 1.5, "07:00", "07:45", "س" * 51, "رياضيات", "معلم أ"]]))
        self.assertTrue(parsed["errors"])
        self.assertGreaterEqual(len(parsed["rows"][0]["errors"]), 2)

    def test_undo_does_not_cascade_delete_new_periods_added_after_import(self):
        parsed = parse_workbook(self.workbook([["الأحد", 1, "07:00", "07:45", "أول أ", "رياضيات", "معلم أ"]]))
        result = apply_import(school=self.school, parsed=parsed)
        day = DaySchedule.objects.get(settings__school=self.school)
        later = Period.objects.create(day=day, index=2, starts_at=time(8), ends_at=time(8, 45))
        with self.assertRaises(ValueError):
            undo_import(school=self.school, undo=result["undo"])
        self.assertTrue(Period.objects.filter(pk=later.pk).exists())
